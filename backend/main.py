import os
import sys
import fitz  # this is okay now — PyMuPDF takes over the name 
import requests
import json
from pptx import Presentation

def extract_text_from_pdf(file_path):
    doc = fitz.open(file_path)
    full_text = ""
    for page in doc:
        full_text += page.get_text()
    return full_text

def call_gemini_api(prompt, api_key):
    url = f"https://generativelanguage.googleapis.com/v1/models/gemini-1.5-flash:generateContent?key={api_key}"
    headers = { "Content-Type": "application/json" }
    body = {
        "contents": [
            {
                "parts": [
                    { "text": prompt }
                ]
            }
        ]
    }

    print("📤 Sending request to Gemini API...")
    res = requests.post(url, headers=headers, data=json.dumps(body))

    print(f"🔍 Gemini API status code: {res.status_code}")
    try:
        res_json = res.json()
        print("🔵 Gemini raw response:")
        print(json.dumps(res_json, indent=2))
        text = res_json['candidates'][0]['content']['parts'][0]['text']
        return json.loads(text)
    except Exception as e:
        print("❌ Error parsing Gemini response:", e)
        print("❌ Full response content:", res.text)
        return []

def generate_ppt(slides, filename="output/presentation.pptx"):
    os.makedirs("output", exist_ok=True)
    prs = Presentation()
    for slide_data in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        content = slide.placeholders[1]

        title.text = slide_data.get("title", "Untitled Slide")
        content.text = "\n".join(slide_data.get("points", []))
    prs.save(filename)

def main():
    file_path = sys.argv[1]
    user_prompt = sys.argv[2]

    # 🔐 REPLACE THIS with your Gemini API key!
    GEMINI_API_KEY = "AIzaSyC1vbVEqNFxn9ssybIMKIOueAi5CZSzNXA"

    raw_text = extract_text_from_pdf(file_path)

    full_prompt = f"""Extract only important insights from the following PDF content and structure it as JSON slides.
Each slide should have a title and 3–5 bullet points. Skip unnecessary details like tables, legal disclaimers, and ads.
Content:\n{raw_text[:8000]}\nUser Prompt: {user_prompt}\n\nRespond ONLY in this JSON format:
[
  {{
    "title": "Slide Title",
    "points": ["point 1", "point 2"]
  }},
  ...
]
"""

    slides = call_gemini_api(full_prompt, GEMINI_API_KEY)

    print("Slide data returned from Gemini:")
    print(slides)

    if not slides:
        print("❌ No slide data returned from Gemini.")
    else:
        print("✅ Slides successfully generated. Writing to .pptx")

    generate_ppt(slides)

if __name__ == "__main__":
    main()
